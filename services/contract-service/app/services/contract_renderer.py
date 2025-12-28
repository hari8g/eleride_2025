"""
Core contract rendering service.
Loads template, replaces placeholders, saves generated contract.
"""
import os
import uuid
from pathlib import Path
from datetime import datetime
from typing import Tuple, Optional
from docx import Document

from app.config import settings
from app.models import ContractRenderRequest
from app.services.placeholder_mapper import model_to_placeholders
from app.utils.docx_utils import replace_placeholders_in_document


class ContractRenderer:
    """Handles contract rendering from template."""
    
    def __init__(self):
        self.template_path = Path(settings.template_dir) / settings.template_filename
        self.generated_dir = Path(settings.generated_dir)
        self.generated_dir.mkdir(exist_ok=True)
    
    def render_contract(
        self, 
        request: ContractRenderRequest,
        output_filename: Optional[str] = None
    ) -> Tuple[str, Path]:
        """
        Render contract from template.
        
        Args:
            request: Contract render request with all data
            output_filename: Optional custom filename (without extension)
        
        Returns:
            Tuple of (filename, file_path)
        """
        if not self.template_path.exists():
            raise FileNotFoundError(f"Template not found: {self.template_path}")
        
        # Load template
        doc = Document(str(self.template_path))
        
        # Map model to placeholders
        placeholders = model_to_placeholders(request)
        
        # Replace placeholders
        replacement_count = replace_placeholders_in_document(doc, placeholders)
        
        if replacement_count == 0:
            # Warning: no replacements made (might be template issue)
            pass
        
        # Generate filename
        # Check explicitly for None or empty string
        if output_filename is None or output_filename == "":
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            unique_id = str(uuid.uuid4())[:8]
            output_filename = f"rider_agreement_{timestamp}_{unique_id}"
        
        output_path = self.generated_dir / f"{output_filename}.docx"
        
        # Save document
        doc.save(str(output_path))
        
        return output_filename, output_path
    
    def get_template_placeholders(self) -> list[str]:
        """
        Inspect template and return all placeholders found.
        Used for template verification.
        """
        if not self.template_path.exists():
            raise FileNotFoundError(f"Template not found: {self.template_path}")
        
        doc = Document(str(self.template_path))
        from app.utils.docx_utils import find_placeholders_in_document
        return find_placeholders_in_document(doc)

